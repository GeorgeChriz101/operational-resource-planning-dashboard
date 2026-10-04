import math
from pathlib import Path

import pandas as pd
from openpyxl.formatting.rule import CellIsRule
from openpyxl.styles import Font, PatternFill


DATA_FILE = "data/workforce_planning_clean.csv"
OUTPUT_FILE = "output/resource_planning_analysis.xlsx"


# -------------------------------------------------
# LOAD CLEAN DATA
# -------------------------------------------------

df = pd.read_csv(DATA_FILE)

df["date"] = pd.to_datetime(df["date"])


# -------------------------------------------------
# OVERALL KPI SUMMARY
# -------------------------------------------------

total_workload = df["workload"].sum()

required_staff = df["required_staff"].sum()

scheduled_staff = df["scheduled_staff"].sum()

overall_staffing_gap = (
    scheduled_staff - required_staff
)

overall_coverage_rate = round(
    (
        scheduled_staff
        / required_staff
    )
    * 100,
    1
)

average_utilization = round(
    df["utilization"].mean(),
    1
)

understaffed_shifts = int(
    (df["staffing_gap"] < 0).sum()
)

high_risk_shifts = int(
    (df["coverage_risk"] == "High").sum()
)


executive_summary = pd.DataFrame(
    [
        {
            "Metric": "Total Workload",
            "Value": total_workload
        },
        {
            "Metric": "Required Staff",
            "Value": required_staff
        },
        {
            "Metric": "Scheduled Staff",
            "Value": scheduled_staff
        },
        {
            "Metric": "Overall Staffing Gap",
            "Value": overall_staffing_gap
        },
        {
            "Metric": "Overall Coverage Rate (%)",
            "Value": overall_coverage_rate
        },
        {
            "Metric": "Average Utilization (%)",
            "Value": average_utilization
        },
        {
            "Metric": "Understaffed Shifts",
            "Value": understaffed_shifts
        },
        {
            "Metric": "High-Risk Shifts",
            "Value": high_risk_shifts
        }
    ]
)


# -------------------------------------------------
# STATION ANALYSIS
# -------------------------------------------------

station_analysis = (
    df.groupby("station")
    .agg(
        total_workload=(
            "workload",
            "sum"
        ),
        required_staff=(
            "required_staff",
            "sum"
        ),
        scheduled_staff=(
            "scheduled_staff",
            "sum"
        ),
        average_utilization=(
            "utilization",
            "mean"
        ),
        understaffed_shifts=(
            "staffing_gap",
            lambda x: (x < 0).sum()
        ),
        high_risk_shifts=(
            "coverage_risk",
            lambda x: (x == "High").sum()
        )
    )
    .reset_index()
)


station_analysis["staffing_gap"] = (
    station_analysis["scheduled_staff"]
    - station_analysis["required_staff"]
)


station_analysis["coverage_rate"] = (
    (
        station_analysis["scheduled_staff"]
        / station_analysis["required_staff"]
    )
    * 100
).round(1)


station_analysis[
    "average_utilization"
] = station_analysis[
    "average_utilization"
].round(1)


# -------------------------------------------------
# SHIFT ANALYSIS
# -------------------------------------------------

shift_analysis = (
    df.groupby("shift")
    .agg(
        total_workload=(
            "workload",
            "sum"
        ),
        required_staff=(
            "required_staff",
            "sum"
        ),
        scheduled_staff=(
            "scheduled_staff",
            "sum"
        ),
        average_utilization=(
            "utilization",
            "mean"
        ),
        understaffed_shifts=(
            "staffing_gap",
            lambda x: (x < 0).sum()
        ),
        high_risk_shifts=(
            "coverage_risk",
            lambda x: (x == "High").sum()
        )
    )
    .reset_index()
)


shift_analysis["staffing_gap"] = (
    shift_analysis["scheduled_staff"]
    - shift_analysis["required_staff"]
)


shift_analysis["coverage_rate"] = (
    (
        shift_analysis["scheduled_staff"]
        / shift_analysis["required_staff"]
    )
    * 100
).round(1)


shift_analysis[
    "average_utilization"
] = shift_analysis[
    "average_utilization"
].round(1)


# -------------------------------------------------
# DAILY TREND
# -------------------------------------------------

daily_trend = (
    df.groupby("date")
    .agg(
        total_workload=(
            "workload",
            "sum"
        ),
        required_staff=(
            "required_staff",
            "sum"
        ),
        scheduled_staff=(
            "scheduled_staff",
            "sum"
        )
    )
    .reset_index()
)


daily_trend["staffing_gap"] = (
    daily_trend["scheduled_staff"]
    - daily_trend["required_staff"]
)


daily_trend["coverage_rate"] = (
    (
        daily_trend["scheduled_staff"]
        / daily_trend["required_staff"]
    )
    * 100
).round(1)


# -------------------------------------------------
# HIGH-RISK SHIFTS
# -------------------------------------------------

high_risk = df[
    df["coverage_risk"] == "High"
].copy()


high_risk = high_risk[
    [
        "date",
        "day_of_week",
        "station",
        "shift",
        "workload",
        "required_staff",
        "scheduled_staff",
        "staffing_gap",
        "coverage_rate",
        "utilization"
    ]
]


high_risk = high_risk.sort_values(
    by=[
        "coverage_rate",
        "workload"
    ],
    ascending=[
        True,
        False
    ]
)


# -------------------------------------------------
# FORECASTING SCENARIOS
# -------------------------------------------------

forecast_rows = []

growth_scenarios = [
    0,
    5,
    10,
    15
]


for growth_rate in growth_scenarios:

    growth_multiplier = (
        1 + growth_rate / 100
    )

    projected_workload = round(
        total_workload
        * growth_multiplier
    )

    projected_required_staff = math.ceil(
        required_staff
        * growth_multiplier
    )

    projected_staffing_gap = (
        scheduled_staff
        - projected_required_staff
    )

    projected_coverage_rate = round(
        (
            scheduled_staff
            / projected_required_staff
        )
        * 100,
        1
    )

    if projected_coverage_rate < 85:
        risk_level = "High"

    elif projected_coverage_rate < 100:
        risk_level = "Medium"

    else:
        risk_level = "Low"


    forecast_rows.append(
        {
            "Workload Growth (%)":
                growth_rate,

            "Projected Workload":
                projected_workload,

            "Projected Required Staff":
                projected_required_staff,

            "Current Scheduled Staff":
                scheduled_staff,

            "Projected Staffing Gap":
                projected_staffing_gap,

            "Projected Coverage Rate (%)":
                projected_coverage_rate,

            "Projected Risk":
                risk_level
        }
    )


forecast_scenarios = pd.DataFrame(
    forecast_rows
)


# -------------------------------------------------
# CREATE OUTPUT DIRECTORY
# -------------------------------------------------

Path("output").mkdir(
    exist_ok=True
)


# -------------------------------------------------
# EXPORT EXCEL WORKBOOK
# -------------------------------------------------

with pd.ExcelWriter(
    OUTPUT_FILE,
    engine="openpyxl"
) as writer:

    executive_summary.to_excel(
        writer,
        sheet_name="Executive Summary",
        index=False
    )

    station_analysis.to_excel(
        writer,
        sheet_name="Station Analysis",
        index=False
    )

    shift_analysis.to_excel(
        writer,
        sheet_name="Shift Analysis",
        index=False
    )

    daily_trend.to_excel(
        writer,
        sheet_name="Daily Trend",
        index=False
    )

    high_risk.to_excel(
        writer,
        sheet_name="High Risk Shifts",
        index=False
    )

    forecast_scenarios.to_excel(
        writer,
        sheet_name="Forecast Scenarios",
        index=False
    )

    df.to_excel(
        writer,
        sheet_name="Clean Data",
        index=False
    )


    workbook = writer.book


    # ---------------------------------------------
    # BASIC FORMATTING
    # ---------------------------------------------

    for worksheet in workbook.worksheets:

        worksheet.freeze_panes = "A2"

        for cell in worksheet[1]:
            cell.font = Font(
                bold=True
            )

        for column_cells in worksheet.columns:

            max_length = 0

            column_letter = (
                column_cells[0]
                .column_letter
            )

            for cell in column_cells:

                value = cell.value

                if value is not None:

                    max_length = max(
                        max_length,
                        len(str(value))
                    )

            worksheet.column_dimensions[
                column_letter
            ].width = min(
                max_length + 2,
                35
            )


    # ---------------------------------------------
    # CONDITIONAL FORMATTING
    # HIGH-RISK STAFFING GAPS
    # ---------------------------------------------

    high_risk_sheet = workbook[
        "High Risk Shifts"
    ]

    red_fill = PatternFill(
        start_color="FFC7CE",
        end_color="FFC7CE",
        fill_type="solid"
    )

    high_risk_sheet.conditional_formatting.add(
        f"H2:H{high_risk_sheet.max_row}",
        CellIsRule(
            operator="lessThan",
            formula=["0"],
            fill=red_fill
        )
    )


print(
    "Excel resource planning report "
    "created successfully."
)

print(
    f"Output file: {OUTPUT_FILE}"
)